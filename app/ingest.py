import os
import io
from pathlib import Path
from typing import List, Dict, Any, Callable, Optional
from concurrent.futures import ThreadPoolExecutor

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

from app.llm import describe_image


SUPPORTED_TEXT_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"}
SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
SUPPORTED_EXTENSIONS = SUPPORTED_TEXT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS

MAX_VISION_CALLS_PER_UPLOAD = 30
MAX_CONCURRENT_VISION_WORKERS = 3


def resize_image_if_needed(image_bytes: bytes, max_edge: int = 1600) -> tuple[bytes, str]:
    """Resize image so its longest edge is at most max_edge px. Returns (bytes, mime_type)."""
    if Image is None:
        return image_bytes, "image/jpeg"

    try:
        img = Image.open(io.BytesIO(image_bytes))
        fmt = img.format or "JPEG"
        mime_type = f"image/{fmt.lower()}"
        if fmt.lower() == "jpg":
            mime_type = "image/jpeg"

        w, h = img.size
        if max(w, h) > max_edge:
            if w >= h:
                new_w = max_edge
                new_h = int(h * (max_edge / w))
            else:
                new_h = max_edge
                new_w = int(w * (max_edge / h))

            img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)

        out_buf = io.BytesIO()
        if img.mode in ("RGBA", "P") and fmt.upper() in ("JPEG", "JPG"):
            img = img.convert("RGB")
        img.save(out_buf, format=fmt)
        return out_buf.getvalue(), mime_type

    except Exception:
        return image_bytes, "image/jpeg"


def extract_image_text(file_path: str | Path) -> List[Dict[str, Any]]:
    """Extract and transcribe text from standalone image file (.jpg, .jpeg, .png, .webp)."""
    path = Path(file_path)
    raw_bytes = path.read_bytes()
    resized_bytes, mime = resize_image_if_needed(raw_bytes)

    transcription = describe_image(resized_bytes, mime_type=mime, mode="FULL")

    return [{
        "text": transcription,
        "source": path.name,
        "page": "photo",
        "is_vision": True
    }]


def extract_text(
    file_path: str | Path,
    progress_callback: Optional[Callable[[str], None]] = None,
    vision_counter: Optional[List[int]] = None
) -> List[Dict[str, Any]]:
    """
    Extract text content from a single file (PDF, PPTX, TXT, MD, Images).
    Returns a list of dicts: [{"text": str, "source": str, "page": int | str, ...}]
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()
    source_name = path.name

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{suffix}' for '{source_name}'. "
            f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if vision_counter is None:
        vision_counter = [0]

    results: List[Dict[str, Any]] = []

    try:
        # Standalone Image Files
        if suffix in SUPPORTED_IMAGE_EXTENSIONS:
            if vision_counter[0] >= MAX_VISION_CALLS_PER_UPLOAD:
                if progress_callback:
                    progress_callback(f"Skipping vision for '{source_name}' (Vision limit {MAX_VISION_CALLS_PER_UPLOAD} reached).")
                return []
            
            vision_counter[0] += 1
            if progress_callback:
                progress_callback(f"Transcribing image '{source_name}'...")
            return extract_image_text(path)

        # PDF Files (Hybrid Text + OCR / Figures)
        elif suffix == ".pdf":
            if PdfReader is None:
                raise ImportError("pypdf is required. Run `pip install pypdf`.")

            reader = PdfReader(path)
            pdfium_doc = None
            if pdfium is not None:
                try:
                    pdfium_doc = pdfium.PdfDocument(path)
                except Exception:
                    pdfium_doc = None

            total_pages = len(reader.pages)

            for idx, page in enumerate(reader.pages, start=1):
                if progress_callback:
                    progress_callback(f"Reading page {idx} of {total_pages} for '{source_name}'...")

                text = (page.extract_text() or "").strip()

                # Check if scanned page (< 30 chars text)
                if len(text) < 30 and pdfium_doc is not None:
                    if vision_counter[0] < MAX_VISION_CALLS_PER_UPLOAD:
                        vision_counter[0] += 1
                        try:
                            pdfium_page = pdfium_doc[idx - 1]
                            image = pdfium_page.render(scale=2).to_pil()
                            buf = io.BytesIO()
                            image.save(buf, format="JPEG")
                            rendered_bytes, mime = resize_image_if_needed(buf.getvalue())

                            ocr_text = describe_image(rendered_bytes, mime_type=mime, mode="FULL")
                            if ocr_text and not ocr_text.startswith("[Vision Error]") and not ocr_text.startswith("[LLM Error]"):
                                results.append({
                                    "text": ocr_text,
                                    "source": source_name,
                                    "page": idx,
                                    "is_vision": True
                                })
                                continue
                        except Exception as e:
                            print(f"[Warning] Failed vision OCR on {source_name} p.{idx}: {e}")

                # Page has embedded images and some text -> transcribe figures
                has_images = len(getattr(page, "images", [])) > 0
                if has_images and pdfium_doc is not None and vision_counter[0] < MAX_VISION_CALLS_PER_UPLOAD:
                    vision_counter[0] += 1
                    try:
                        pdfium_page = pdfium_doc[idx - 1]
                        image = pdfium_page.render(scale=2).to_pil()
                        buf = io.BytesIO()
                        image.save(buf, format="JPEG")
                        rendered_bytes, mime = resize_image_if_needed(buf.getvalue())

                        fig_desc = describe_image(rendered_bytes, mime_type=mime, mode="FIGURES")
                        if fig_desc and not fig_desc.startswith("[Vision Error]") and not fig_desc.startswith("[LLM Error]"):
                            combined = f"{text}\n\n[Diagram/Figure Description]:\n{fig_desc}".strip()
                            results.append({
                                "text": combined,
                                "source": source_name,
                                "page": idx,
                                "is_vision": True
                            })
                            continue
                    except Exception as e:
                        print(f"[Warning] Failed vision figure description on {source_name} p.{idx}: {e}")

                if text:
                    results.append({"text": text, "source": source_name, "page": idx})

        # PPTX Files (Slide Text + Image-only slide OCR fallback)
        elif suffix == ".pptx":
            if Presentation is None:
                raise ImportError("python-pptx is required. Run `pip install python-pptx`.")

            prs = Presentation(path)
            total_slides = len(prs.slides)

            for idx, slide in enumerate(prs.slides, start=1):
                if progress_callback and idx % 5 == 1:
                    progress_callback(f"Reading slide {idx} of {total_slides} for '{source_name}'...")

                slide_texts = []
                picture_blobs = []

                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame:
                        slide_texts.append(shape.text_frame.text)
                    if hasattr(shape, "image") and shape.image:
                        picture_blobs.append(shape.image.blob)

                text_content = "\n".join(slide_texts).strip()

                # Only run vision OCR if slide has no text content (< 10 chars) and has pictures
                if len(text_content) < 10 and picture_blobs and vision_counter[0] < MAX_VISION_CALLS_PER_UPLOAD:
                    vision_counter[0] += 1
                    try:
                        pic_bytes, mime = resize_image_if_needed(picture_blobs[0])
                        fig_desc = describe_image(pic_bytes, mime_type=mime, mode="FULL")
                        if fig_desc and not fig_desc.startswith("[Vision Error]") and not fig_desc.startswith("[LLM Error]"):
                            text_content = fig_desc.strip()
                    except Exception as e:
                        print(f"[Warning] Failed figure vision on slide {idx}: {e}")

                if text_content:
                    results.append({"text": text_content, "source": source_name, "page": idx})

        # Plain Text / Markdown
        elif suffix in (".txt", ".md"):
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = path.read_text(encoding="latin-1")

            cleaned = content.strip()
            if cleaned:
                results.append({"text": cleaned, "source": source_name, "page": 1})

    except Exception as e:
        if isinstance(e, (ValueError, ImportError, FileNotFoundError)):
            raise e
        raise ValueError(f"Failed to extract text from '{source_name}': {str(e)}") from e

    return results


def ingest_directory(
    dir_path: str | Path,
    progress_callback: Optional[Callable[[str], None]] = None
) -> List[Dict[str, Any]]:
    """Recursively extract text & vision transcriptions from supported files in directory."""
    directory = Path(dir_path)
    if not directory.is_dir():
        raise NotADirectoryError(f"Directory not found: {dir_path}")

    all_extracted: List[Dict[str, Any]] = []
    vision_counter = [0]

    for root, _, files in os.walk(directory):
        for file in sorted(files):
            file_path = Path(root) / file
            if file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
                try:
                    extracted = extract_text(file_path, progress_callback, vision_counter)
                    all_extracted.extend(extracted)
                except Exception as err:
                    print(f"[Warning] Skipping {file_path.name}: {err}")

    return all_extracted
